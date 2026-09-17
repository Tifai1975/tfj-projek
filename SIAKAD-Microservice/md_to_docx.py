#!/usr/bin/env python3
"""Convert markdown PRD to Word .docx with proper styling"""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def parse_markdown_to_docx(md_content, output_path):
    """Parse markdown and create Word document"""
    doc = Document()
    
    # Set up styles
    styles = doc.styles
    
    # Custom heading styles with colors
    try:
        h1_style = styles['Heading 1']
        h1_style.font.size = Pt(18)
        h1_style.font.color.rgb = RGBColor(44, 62, 80)  # #2c3e50
    except:
        pass
    
    try:
        h2_style = styles['Heading 2']
        h2_style.font.size = Pt(14)
        h2_style.font.color.rgb = RGBColor(52, 73, 94)  # #34495e
    except:
        pass
    
    try:
        h3_style = styles['Heading 3']
        h3_style.font.size = Pt(12)
        h3_style.font.color.rgb = RGBColor(127, 140, 141)  # #7f8c8d
    except:
        pass
    
    lines = md_content.split('\n')
    i = 0
    in_table = False
    table_rows = []
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Skip empty lines
        if not line:
            if not in_table and not in_code_block:
                doc.add_paragraph()  # Add spacing
            i += 1
            continue
        
        # Code blocks
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
                i += 1
                continue
            else:
                # End of code block
                if code_lines:
                    code_para = doc.add_paragraph('\n'.join(code_lines))
                    code_para.style = 'No Spacing'
                    for run in code_para.runs:
                        run.font.name = 'Courier New'
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(199, 37, 78)  # #c7254e
                in_code_block = False
                code_lines = []
                i += 1
                continue
        
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
        
        # Title (first # heading)
        if line.startswith('# ') and i < 5:
            text = line[2:].strip()
            title = doc.add_heading(text, level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        
        # Headings
        if line.startswith('#### '):
            text = line[5:].strip()
            doc.add_heading(text, level=4)
            i += 1
            continue
        elif line.startswith('### '):
            text = line[4:].strip()
            doc.add_heading(text, level=3)
            i += 1
            continue
        elif line.startswith('## '):
            text = line[3:].strip()
            doc.add_heading(text, level=2)
            i += 1
            continue
        elif line.startswith('# '):
            text = line[2:].strip()
            doc.add_heading(text, level=1)
            i += 1
            continue
        
        # Horizontal rule
        if line.startswith('---'):
            doc.add_paragraph('_' * 50)
            i += 1
            continue
        
        # Tables
        if '|' in line and not in_table:
            in_table = True
            table_rows = []
        
        if in_table:
            if '|' in line:
                # Parse table row
                cells = [c.strip() for c in line.split('|')]
                cells = [c for c in cells if c]  # Remove empty
                
                # Skip separator row
                if cells and not all(set(c) <= set('-: ') for c in cells):
                    table_rows.append(cells)
                
                i += 1
                continue
            else:
                # End of table
                if table_rows and len(table_rows) > 1:
                    # Create table
                    num_cols = len(table_rows[0])
                    table = doc.add_table(rows=len(table_rows), cols=num_cols)
                    table.style = 'Light Grid Accent 1'
                    
                    for row_idx, row_data in enumerate(table_rows):
                        for col_idx, cell_text in enumerate(row_data):
                            if col_idx < num_cols:
                                cell = table.rows[row_idx].cells[col_idx]
                                cell.text = cell_text
                                # Bold header row
                                if row_idx == 0:
                                    for paragraph in cell.paragraphs:
                                        for run in paragraph.runs:
                                            run.bold = True
                    
                    doc.add_paragraph()  # Spacing after table
                
                in_table = False
                table_rows = []
                continue
        
        # Bullet lists
        if line.startswith('- ') or line.startswith('* '):
            text = line[2:].strip()
            # Remove markdown formatting
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)  # Bold
            text = re.sub(r'\*(.+?)\*', r'\1', text)  # Italic
            text = re.sub(r'`(.+?)`', r'\1', text)  # Code
            
            para = doc.add_paragraph(text, style='List Bullet')
            i += 1
            continue
        
        # Numbered lists (if starts with number.)
        if re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line).strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            text = re.sub(r'`(.+?)`', r'\1', text)
            
            para = doc.add_paragraph(text, style='List Number')
            i += 1
            continue
        
        # Bold metadata lines
        if line.startswith('**') and '**:' in line:
            para = doc.add_paragraph()
            # Split into bold and normal parts
            parts = line.split('**:')
            if len(parts) == 2:
                bold_text = parts[0].replace('**', '')
                normal_text = parts[1]
                
                run_bold = para.add_run(bold_text + ':')
                run_bold.bold = True
                para.add_run(normal_text)
            else:
                para.add_run(line)
            i += 1
            continue
        
        # Regular paragraphs
        text = line
        
        # Check if contains inline formatting
        if '**' in text or '*' in text or '`' in text:
            para = doc.add_paragraph()
            
            # Simple inline formatting parser
            parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = para.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
                    run = para.add_run(part[1:-1])
                    run.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    run = para.add_run(part[1:-1])
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(199, 37, 78)
                else:
                    para.add_run(part)
        else:
            doc.add_paragraph(text)
        
        i += 1
    
    # Save document
    doc.save(output_path)
    print(f"✅ Word document created: {output_path}")

def main():
    input_file = '/Users/tfj/Documents/Projek Hermes/SIAKAD-Microservice/PRD-SIAKAD.md'
    output_file = '/Users/tfj/Documents/Projek Hermes/SIAKAD-Microservice/PRD-SIAKAD.docx'
    
    # Read markdown
    with open(input_file, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # Convert to docx
    parse_markdown_to_docx(md_content, output_file)

if __name__ == '__main__':
    main()
